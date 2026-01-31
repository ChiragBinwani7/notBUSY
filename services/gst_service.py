"""GST service for generating reports."""

from sqlalchemy.orm import Session
from sqlalchemy import and_, extract
from models.gst import GSTTransaction
from core.config import GSTTransactionType
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import os

class GSTService:
    """Service for GST reporting."""
    
    @staticmethod
    def get_monthly_sales_gst(db: Session, year: int, month: int):
        """Get GSTR-1 (sales) for a specific month."""
        transactions = db.query(GSTTransaction).filter(
            and_(
                GSTTransaction.transaction_type == GSTTransactionType.SALE.value,
                extract('year', GSTTransaction.transaction_date) == year,
                extract('month', GSTTransaction.transaction_date) == month
            )
        ).all()
        
        return transactions
    
    @staticmethod
    def get_monthly_purchase_gst(db: Session, year: int, month: int):
        """Get purchase GST with separate sections for grey purchases and mill job work."""
        grey_purchases = db.query(GSTTransaction).filter(
            and_(
                GSTTransaction.transaction_type == GSTTransactionType.GREY_PURCHASE.value,
                extract('year', GSTTransaction.transaction_date) == year,
                extract('month', GSTTransaction.transaction_date) == month
            )
        ).all()
        
        mill_jobs = db.query(GSTTransaction).filter(
            and_(
                GSTTransaction.transaction_type == GSTTransactionType.MILL_JOB.value,
                extract('year', GSTTransaction.transaction_date) == year,
                extract('month', GSTTransaction.transaction_date) == month
            )
        ).all()
        
        return {
            "grey_purchases": grey_purchases,
            "mill_jobs": mill_jobs
        }
    
    @staticmethod
    def get_gstr3b_summary(db: Session, year: int, month: int):
        """Get GSTR-3B summary for a month."""
        # Sales GST (output)
        sales = GSTService.get_monthly_sales_gst(db, year, month)
        total_sales_gst = sum(t.total_gst for t in sales)
        total_sales_cgst = sum(t.cgst_amount for t in sales)
        total_sales_sgst = sum(t.sgst_amount for t in sales)
        total_sales_igst = sum(t.igst_amount for t in sales)
        
        # Purchase GST (input)
        purchases = GSTService.get_monthly_purchase_gst(db, year, month)
        all_purchases = purchases["grey_purchases"] + purchases["mill_jobs"]
        total_purchase_gst = sum(t.total_gst for t in all_purchases)
        total_purchase_cgst = sum(t.cgst_amount for t in all_purchases)
        total_purchase_sgst = sum(t.sgst_amount for t in all_purchases)
        total_purchase_igst = sum(t.igst_amount for t in all_purchases)
        
        # Net GST payable
        net_gst_payable = total_sales_gst - total_purchase_gst
        
        return {
            "sales": {
                "total_gst": total_sales_gst,
                "cgst": total_sales_cgst,
                "sgst": total_sales_sgst,
                "igst": total_sales_igst,
                "count": len(sales)
            },
            "purchases": {
                "total_gst": total_purchase_gst,
                "cgst": total_purchase_cgst,
                "sgst": total_purchase_sgst,
                "igst": total_purchase_igst,
                "grey_count": len(purchases["grey_purchases"]),
                "mill_count": len(purchases["mill_jobs"])
            },
            "net_gst_payable": net_gst_payable
        }
    
    @staticmethod
    def export_to_excel(db: Session, year: int, month: int, export_dir: str = "exports"):
        """Export GST reports to Excel."""
        if not os.path.exists(export_dir):
            os.makedirs(export_dir)
        
        filename = f"GST_Report_{year}_{month:02d}.xlsx"
        filepath = os.path.join(export_dir, filename)
        
        wb = Workbook()
        
        # GSTR-1 Sheet (Sales)
        ws1 = wb.active
        ws1.title = "GSTR-1 Sales"
        ws1.append(["Date", "Customer", "State", "Taxable Amount", "CGST", "SGST", "IGST", "Total GST"])
        
        sales = GSTService.get_monthly_sales_gst(db, year, month)
        for t in sales:
            ws1.append([
                t.transaction_date.strftime("%Y-%m-%d"),
                t.party_name,
                t.party_state,
                t.taxable_amount,
                t.cgst_amount,
                t.sgst_amount,
                t.igst_amount,
                t.total_gst
            ])
        
        # Purchase GST Sheet
        ws2 = wb.create_sheet("Purchase GST")
        ws2.append(["Type", "Date", "Party", "State", "Taxable Amount", "CGST", "SGST", "IGST", "Total GST"])
        
        purchases = GSTService.get_monthly_purchase_gst(db, year, month)
        
        # Grey purchases section
        for t in purchases["grey_purchases"]:
            ws2.append([
                "Grey Purchase",
                t.transaction_date.strftime("%Y-%m-%d"),
                t.party_name,
                t.party_state,
                t.taxable_amount,
                t.cgst_amount,
                t.sgst_amount,
                t.igst_amount,
                t.total_gst
            ])
        
        # Mill job work section
        for t in purchases["mill_jobs"]:
            ws2.append([
                "Mill Job Work",
                t.transaction_date.strftime("%Y-%m-%d"),
                t.party_name,
                t.party_state,
                t.taxable_amount,
                t.cgst_amount,
                t.sgst_amount,
                t.igst_amount,
                t.total_gst
            ])
        
        # GSTR-3B Summary Sheet
        ws3 = wb.create_sheet("GSTR-3B Summary")
        summary = GSTService.get_gstr3b_summary(db, year, month)
        
        ws3.append(["GSTR-3B Summary", f"{year}-{month:02d}"])
        ws3.append([])
        ws3.append(["Output GST (Sales)"])
        ws3.append(["CGST", summary["sales"]["cgst"]])
        ws3.append(["SGST", summary["sales"]["sgst"]])
        ws3.append(["IGST", summary["sales"]["igst"]])
        ws3.append(["Total Sales GST", summary["sales"]["total_gst"]])
        ws3.append([])
        ws3.append(["Input GST (Purchases)"])
        ws3.append(["CGST", summary["purchases"]["cgst"]])
        ws3.append(["SGST", summary["purchases"]["sgst"]])
        ws3.append(["IGST", summary["purchases"]["igst"]])
        ws3.append(["Total Purchase GST", summary["purchases"]["total_gst"]])
        ws3.append([])
        ws3.append(["Net GST Payable", summary["net_gst_payable"]])
        
        wb.save(filepath)
        return filepath
